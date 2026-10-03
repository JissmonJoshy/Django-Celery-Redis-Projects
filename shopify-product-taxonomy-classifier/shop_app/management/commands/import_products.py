from django.core.management.base import BaseCommand
from django.conf import settings
import openpyxl
import os

from shop_app.models import Product, ProductImage


class Command(BaseCommand):
    help = "Import products from Product_List.xlsx into the database"

    def handle(self, *args, **options):
        file_path = os.path.join(settings.BASE_DIR, "Product_List.xlsx")

        if not os.path.exists(file_path):
            self.stderr.write(f"File not found: {file_path}")
            return

        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        # Note: the actual Excel header is "Product Description " (trailing space)
        col = {name: idx for idx, name in enumerate(headers)}

        created_count = 0
        skipped_count = 0
        error_count = 0

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                def get(col_name, default=""):
                    idx = col.get(col_name)
                    if idx is None:
                        return default
                    val = row[idx]
                    return val if val is not None else default

                product_number = get("Product Number")
                if not product_number:
                    skipped_count += 1
                    continue  # can't import a row with no product number

                product, was_created = Product.objects.update_or_create(
                    product_number=str(product_number),
                    defaults={
                        "model_number": get("Model Number"),
                        "product_category": get("Product Category"),
                        "product_sub_category": get("Product Sub Category"),
                        "collection_name": get("Collection Name"),
                        "color_collection": get("Color Collection"),
                        "product_color": get("Product Color"),
                        "product_name": get("Product Name"),
                        "description": get("Product Description "),  # trailing space in Excel header
                        "bullets": get("Bullets"),
                        "set_includes": get("Set Includes"),
                        "product_weight": str(get("Product Weight")),
                        "materials": get("Materials"),
                        "product_dimensions": str(get("Product Dimensions")),
                        "assembly_required": str(get("Assembly Required")),
                        "is_a_set": str(get("Is a Set")),
                        "stackable": str(get("Stackable")),
                        "country_of_origin": get("Country Of Origin"),
                        "item_cost": get("Item Cost", None),
                        "map_price": get("MAP", None),
                        "msrp": get("MSRP", None),
                        "shipping_method": get("Shipping Method"),
                        "total_box_count": get("Total Box Count", None),
                        "pallet_count": get("Pallet Count", None),
                        "shipping_weight": str(get("Shipping Weight")),
                        "total_cbm": str(get("Total CBM")),
                        "package_dimensions": str(get("Package Dimensions")),
                        "product_url": get("Product URL"),
                    }
                )

                # Clear old images on re-import, then add current ones
                product.images.all().delete()
                for i in range(1, 21):
                    url = get(f"Image {i}")
                    if url:
                        ProductImage.objects.create(
                            product=product, image_url=str(url), position=i
                        )

                created_count += 1

            except Exception as e:
                error_count += 1
                self.stderr.write(f"Row {row_num} failed: {e}")
                continue  # never let one bad row stop the whole import

        self.stdout.write(self.style.SUCCESS(
            f"Done. Imported/updated: {created_count}, skipped (no product number): {skipped_count}, errors: {error_count}"
        ))