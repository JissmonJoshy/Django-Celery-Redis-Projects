from django.shortcuts import render, redirect,get_object_or_404
from django.contrib.auth import authenticate, login as auth_login
from django.contrib import messages
from django.contrib.auth.decorators import login_required

from .models import *
import openpyxl
from django.conf import settings as dj_settings
import os

from django.core.paginator import Paginator
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from .serializers import ClassificationResultSerializer
from rest_framework.pagination import PageNumberPagination



# Create your views here.
def index(request):
    return render(request, 'index.html')




def view_login(request):

    if request.method == 'POST':

        username = request.POST['username']
        password = request.POST['password']

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None and user.is_active:

            auth_login(request, user)

            if user.username == "admin":

                messages.success(request, "Admin login successful")
                return redirect('admin_dashboard')

            elif user.usertype == "reviewer":

                request.session['uid'] = user.id

                messages.success(request, "Reviewer login successful")
                return redirect('reviewer_dashboard')

            else:

                messages.error(request, "User type not recognized.")
                return render(request, 'login.html')

        else:

            messages.error(request, "Invalid username or password.")
            return render(request, 'login.html')

    return render(request, 'login.html')





def reviewer_register(request):

    if request.method == 'POST':

        name = request.POST['name']
        address = request.POST['address']
        email = request.POST['email']
        phone = request.POST['phone']
        username = request.POST['username']
        password = request.POST['password']
        confirm_password = request.POST['confirm_password']
        image = request.FILES.get('image')

        # Check password
        if password != confirm_password:
            messages.error(request, "Passwords do not match.")
            return render(request, 'reviewer/reviewer_register.html')

        # Check username
        if Login.objects.filter(username=username).exists():
            messages.error(request, "Username already exists.")
            return render(request, 'reviewer/reviewer_register.html')

        # Check email
        if Login.objects.filter(email=email).exists():
            messages.error(request, "Email already exists.")
            return render(request, 'reviewer/reviewer_register.html')

        # Check phone
        if Reviewer.objects.filter(phone=phone).exists():
            messages.error(request, "Phone number already exists.")
            return render(request, 'reviewer/reviewer_register.html')

        # Create Login
        login = Login.objects.create_user(
            username=username,
            password=password,
            email=email,
            is_active=True
        )

        # Store password for display if your project requires it
        login.usertype = "reviewer" 
        login.viewpassword = password
        login.save()

        # Create Reviewer
        Reviewer.objects.create(
            login=login,
            name=name,
            address=address,
            email=email,
            phone=phone,
            image=image
        )

        messages.success(
            request,
            "Reviewer registration successful. Please login."
        )

        return redirect('login')

    return render(request, 'reviewer_register.html')


@login_required
def admin_dashboard(request):
    return render(request, 'admin/admin_dashboard.html')

@login_required
def reviewer_dashboard(request):
    return render(request, 'reviewer/reviewer_dashboard.html')






@login_required
def product_list(request):
    status_filter = request.GET.get('status', '')

    qs = ClassificationResult.objects.select_related('product', 'predicted_category').all()
    if status_filter:
        qs = qs.filter(status=status_filter)
    qs = qs.order_by('confidence')  # lowest confidence (needs attention most) first

    paginator = Paginator(qs, 25)
    page_obj = paginator.get_page(request.GET.get('page'))

    template = 'admin/product_list.html' if request.user.usertype == '' else 'reviewer/product_list.html'
    return render(request, template, {'page_obj': page_obj, 'status_filter': status_filter})


@login_required
def product_detail(request, result_id):
    result = get_object_or_404(ClassificationResult, id=result_id)
    categories = Category.objects.all().order_by('full_path')

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'approve':
            result.status = 'approved'
            result.save()
            messages.success(request, "Classification approved.")
        elif action == 'correct':
            new_cat_id = request.POST.get('category')
            if new_cat_id:
                result.predicted_category_id = new_cat_id
                result.status = 'approved'
                result.confidence = 1.0
                result.is_low_confidence = False
                result.save()
                messages.success(request, "Category corrected and approved.")
        return redirect('product_detail', result_id=result.id)

    template = 'admin/product_detail.html' if request.user.usertype == '' else 'reviewer/product_detail.html'
    return render(request, template, {
        'result': result,
        'categories': categories,
        'alternatives': result.get_alternatives(),
    })





@api_view(['GET'])
@permission_classes([IsAuthenticated])
def api_classification_list(request):
    status_filter = request.GET.get('status')
    qs = ClassificationResult.objects.select_related('product', 'predicted_category').prefetch_related('attribute_values')
    if status_filter:
        qs = qs.filter(status=status_filter)
    qs = qs.order_by('confidence')

    paginator = PageNumberPagination()
    page = paginator.paginate_queryset(qs, request)
    serializer = ClassificationResultSerializer(page, many=True)
    return paginator.get_paginated_response(serializer.data)



@api_view(['GET'])
@permission_classes([IsAuthenticated])
def api_classification_detail(request, pk):
    result = get_object_or_404(ClassificationResult, pk=pk)
    serializer = ClassificationResultSerializer(result)
    return Response(serializer.data)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def api_classification_approve(request, pk):
    result = get_object_or_404(ClassificationResult, pk=pk)
    result.status = 'approved'
    result.save()
    return Response({'status': 'approved', 'id': result.id})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def api_classification_update(request, pk):
    result = get_object_or_404(ClassificationResult, pk=pk)
    category_id = request.data.get('category_id')
    if category_id:
        result.predicted_category_id = category_id
        result.status = 'approved'
        result.confidence = 1.0
        result.is_low_confidence = False
        result.save()
        return Response({'status': 'updated', 'category_id': category_id})
    return Response({'error': 'category_id required'}, status=400)








@login_required
def import_products_view(request):
    if request.user.usertype != '':
        messages.error(request, "Only admin can import products.")
        return redirect('admin_dashboard')

    if request.method == 'POST':
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            messages.error(request, "Please choose a file.")
            return render(request, 'admin/import_products.html')

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            ws = wb.active
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            col = {name: idx for idx, name in enumerate(headers)}

            created, skipped, errors = 0, 0, 0
            error_samples = []  # keep the first few actual error messages so we can see what's wrong

            def get(row, name, default=""):
                idx = col.get(name)
                if idx is None:
                    return default
                val = row[idx]
                return val if val is not None else default

            rows = list(ws.iter_rows(min_row=2, values_only=True))

            for row in rows:
                try:
                    product_number = get(row, "Product Number")
                    if not product_number:
                        skipped += 1
                        continue

                    product, _ = Product.objects.update_or_create(
                        product_number=str(product_number)[:100],
                        defaults={
                            "product_name": str(get(row, "Product Name"))[:500],
                            "description": str(get(row, "Product Description ")),
                            "product_category": str(get(row, "Product Category"))[:200],
                            "product_sub_category": str(get(row, "Product Sub Category"))[:200],
                            "materials": str(get(row, "Materials"))[:300],
                            "product_color": str(get(row, "Product Color"))[:200],
                            "assembly_required": str(get(row, "Assembly Required"))[:10],
                            "product_url": str(get(row, "Product URL"))[:500],
                        }
                    )

                    product.images.all().delete()
                    image_objs = []
                    for i in range(1, 21):
                        url = get(row, f"Image {i}")
                        if url:
                            image_objs.append(ProductImage(product=product, image_url=str(url)[:500], position=i))
                    if image_objs:
                        ProductImage.objects.bulk_create(image_objs)

                    created += 1

                except Exception as e:
                    errors += 1
                    if len(error_samples) < 5:
                        error_samples.append(f"{get(row, 'Product Number', 'unknown')}: {e}")
                    continue

            msg = f"Import complete. Created/updated: {created}, skipped: {skipped}, errors: {errors}"
            if error_samples:
                msg += " | Sample errors: " + " || ".join(error_samples)
            messages.success(request, msg)
            return redirect('admin_dashboard')

        except Exception as e:
            messages.error(request, f"Could not read file: {e}")
            return render(request, 'admin/import_products.html')

    return render(request, 'admin/import_products.html')



from shop_app.tasks import classify_products_batch_task

@login_required
def run_classification_view(request):
    classify_products_batch_task.delay()  # queues it, returns instantly, worker picks it up
    messages.success(request, "Classification started in the background. Check the review page shortly.")
    return redirect('admin_dashboard')